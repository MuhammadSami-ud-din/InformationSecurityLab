import sys
import openpyxl
from Crypto.Cipher import AES
from Crypto.Util.Padding import unpad

HEX = "0123456789abcdef"


def caesar_unshift(hex_str, s):
    """Undo a Caesar shift of s on hex digits: (nibble - s) mod 16."""
    return "".join(HEX[(int(c, 16) - s) % 16] for c in hex_str.lower())


def load_keys(path):
    ws = openpyxl.load_workbook(path).active
    return [str(r[0]).strip() for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


def load_student(path):
    ws = openpyxl.load_workbook(path).active
    row = next(ws.iter_rows(min_row=2, values_only=True))
    return [bytes.fromhex(str(v).strip()) for v in row]


def try_decrypt(ct, key):
    """AES-256-ECB decrypt + PKCS#7 unpad. Returns text or None."""
    try:
        pt = unpad(AES.new(key, AES.MODE_ECB).decrypt(ct), 16).decode("utf-8")
        return pt if pt.isprintable() else None
    except (ValueError, UnicodeDecodeError):
        return None


def main(student_no):
    keys = load_keys("ENC_KEYS.xlsx")
    cts = load_student(f"ENC_STUDENT_DATA/{student_no}.xlsx")
    print(f"[+] Keys loaded: {len(keys)} | Ciphertext lengths (bytes): {[len(c) for c in cts]}")

    with open("candidates.txt", "w") as f:
        for ki, k in enumerate(keys, 1):
            for s in range(16):
                f.write(f"key#{ki} shift={s:2d} {caesar_unshift(k, s)}\n")
    print(f"[+] {len(keys) * 16} candidate keys saved to candidates.txt")


    for ki, enc_key in enumerate(keys, 1):
        for s in range(16):
            cand = caesar_unshift(enc_key, s)
            key = bytes.fromhex(cand)
            assert len(key) == 32
            res = [try_decrypt(ct, key) for ct in cts]
            if all(r is not None for r in res):
                print(f"\n[MATCH] key row #{ki}, shift = {s}")
                print(f"  Encrypted key : {enc_key}")
                print(f"  AES-256 key   : {cand}")
                print(f"  Plaintext 1   : {res[0]}")
                print(f"  Plaintext 2   : {res[1]}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "28")